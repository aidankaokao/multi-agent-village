"""讀檔：xlsx / csv / txt / pdf → 統一的「文件內容」（文字＋表格統計）。全部由程式處理，不靠 LLM。

上傳檔當不可信資料：份數、大小、列數、頁數都有上限；xlsx 先檢查解壓後大小（防壓縮炸彈）；只在記憶體處理，不落地。
數字欄的總和 / 平均 / 最大 / 最小由程式算好，報告裡的數字以這裡為準。
"""

import csv
import io
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import PurePath

from app.config import Limits

SUPPORTED = {".xlsx": "xlsx", ".csv": "csv", ".txt": "txt", ".pdf": "pdf"}
KIND_LABEL = {"xlsx": "Excel", "csv": "CSV", "txt": "文字檔", "pdf": "PDF"}
SAMPLE_ROWS = 20  # 每個表格給 LLM 看前幾列
TOP_VALUES = 5  # 文字欄列出最常見的幾個值
GROUP_MAX_KEYS = 12  # 文字欄的不同值不超過這個數，才拿來做分組小計（例如部門、品項）
GROUP_BY_COLS = 2  # 最多用幾個文字欄分組
GROUP_SUM_COLS = 8  # 每組最多統計幾個數字欄


class ReadError(ValueError):
    """這份檔案不能讀（格式不支援、太大、掃描檔、加密…）；訊息直接給使用者看。"""


@dataclass
class Column:
    name: str
    filled: int = 0  # 非空值的數量
    numeric: dict | None = None  # {count, sum, mean, min, max}
    dates: dict | None = None  # {min, max}
    top: list[tuple[str, int]] = field(default_factory=list)  # 文字欄最常見的值

    def to_dict(self) -> dict:
        return {"name": self.name, "filled": self.filled, "numeric": self.numeric, "dates": self.dates, "top": self.top}


@dataclass
class Table:
    name: str  # 工作表名稱（CSV 就是檔名）
    columns: list[Column]
    rows: int  # 實際讀到的資料列數（不含標題列）
    sample: list[list[str]]
    truncated: bool = False  # 超過列數上限，後面沒讀
    # 分組小計：[{by: 欄名, sums: [欄名…], groups: [{key, count, sums: [合計…], means: [平均…]}]}]
    # （程式算好，報告裡的小計以這裡為準；客單價、比率這類欄位看平均，合計沒有意義）
    groups: list[dict] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {"name": self.name, "rows": self.rows, "truncated": self.truncated,
                "columns": [c.to_dict() for c in self.columns], "sample": self.sample, "groups": self.groups}


@dataclass
class Document:
    name: str
    kind: str  # xlsx / csv / txt / pdf
    size: int
    text: str = ""  # txt、pdf 的內文（pdf 每頁前面有「〔第 N 頁〕」）
    tables: list[Table] = field(default_factory=list)
    pages: int | None = None
    encoding: str | None = None
    notes: list[str] = field(default_factory=list)  # 讀檔時的提醒（例如超過上限沒讀完）

    @property
    def label(self) -> str:
        return KIND_LABEL[self.kind]

    def info(self) -> str:
        """一行基本資料（附錄、提示詞用）。"""
        parts = [self.label, _size(self.size)]
        if self.tables:
            parts.append(f"{len(self.tables)} 個表格、共 {sum(t.rows for t in self.tables):,} 列")
        if self.pages is not None:
            parts.append(f"{self.pages} 頁")
        if self.text:
            parts.append(f"約 {len(self.text):,} 字")
        if self.encoding:
            parts.append(f"編碼 {self.encoding}")
        return "、".join(parts)

    def to_dict(self) -> dict:
        return {"name": self.name, "kind": self.kind, "size": self.size, "pages": self.pages, "encoding": self.encoding,
                "chars": len(self.text), "tables": [t.to_dict() for t in self.tables], "notes": self.notes,
                "info": self.info()}


def _size(n: int) -> str:
    return f"{n / 1024 / 1024:.1f} MB" if n >= 1024 * 1024 else f"{max(1, round(n / 1024))} KB"


def kind_of(name: str) -> str:
    ext = PurePath(name).suffix.lower()
    if ext not in SUPPORTED:
        hint = "（舊版 .xls 請另存成 .xlsx）" if ext == ".xls" else ""
        raise ReadError(f"「{name}」的格式不支援{hint}；可以上傳 {'、'.join(SUPPORTED)}")
    return SUPPORTED[ext]


def check_batch(files: list[tuple[str, bytes]], limits: Limits) -> None:
    """整批的份數與大小（讀檔之前先擋）。"""
    if not files:
        raise ReadError("沒有收到檔案")
    if len(files) > limits.max_files:
        raise ReadError(f"最多 {limits.max_files} 份檔案，這次收到 {len(files)} 份")
    for name, data in files:
        kind_of(name)
        if len(data) > limits.max_file_bytes:
            raise ReadError(f"「{name}」有 {_size(len(data))}，超過每份 {_size(limits.max_file_bytes)} 的上限")
        if not data:
            raise ReadError(f"「{name}」是空檔案")
    total = sum(len(d) for _, d in files)
    if total > limits.max_total_bytes:
        raise ReadError(f"檔案合計 {_size(total)}，超過 {_size(limits.max_total_bytes)} 的上限")


def read(name: str, data: bytes, limits: Limits) -> Document:
    kind = kind_of(name)
    doc = Document(name=name, kind=kind, size=len(data))
    {"xlsx": _xlsx, "csv": _csv, "txt": _txt, "pdf": _pdf}[kind](doc, data, limits)
    return doc


# ---------------------------------------------------------------- 文字編碼

# 台灣常見：UTF-8（含 BOM）、Big5 / cp950（舊版 Excel 另存的 CSV、舊系統匯出的 txt）
_ENCODINGS = ["utf-8-sig", "cp950", "big5hkscs", "gb18030"]


def decode(data: bytes) -> tuple[str, str]:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), "UTF-16"
    for enc in _ENCODINGS:
        try:
            text = data.decode(enc)
        except UnicodeDecodeError:
            continue
        return text, {"utf-8-sig": "UTF-8", "cp950": "Big5"}.get(enc, enc.upper())
    return data.decode("utf-8", errors="replace"), "UTF-8（有無法辨識的字元）"


# ---------------------------------------------------------------- 表格統計


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M") if (v.hour or v.minute) else v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _number(v) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace(",", "").replace("$", "").replace("NT", "")
        if not s or len(s) > 30:
            return None
        try:
            return float(s)
        except ValueError:
            return None
    return None


def _round(x: float) -> float | int:
    return int(x) if float(x).is_integer() else round(x, 4)


class _ColumnStats:
    def __init__(self, name: str):
        self.name = name
        self.filled = 0
        self.nums: list[float] = []
        self.dates: list[date] = []
        self.texts: Counter[str] = Counter()

    @property
    def distinct(self) -> int:
        return len(self.texts)

    def add(self, v) -> None:
        if v is None or (isinstance(v, str) and not v.strip()):
            return
        self.filled += 1
        if isinstance(v, (datetime, date)):
            self.dates.append(v.date() if isinstance(v, datetime) else v)
        elif (n := _number(v)) is not None:
            self.nums.append(n)
        else:
            self.texts[_cell(v)[:40]] += 1

    def result(self) -> Column:
        col = Column(self.name, filled=self.filled)
        if self.nums and len(self.nums) >= self.filled * 0.6:  # 大部分是數字才算數字欄
            total = sum(self.nums)
            col.numeric = {"count": len(self.nums), "sum": _round(total), "mean": _round(total / len(self.nums)),
                           "min": _round(min(self.nums)), "max": _round(max(self.nums))}
        elif self.dates and len(self.dates) >= self.filled * 0.6:
            col.dates = {"min": min(self.dates).isoformat(), "max": max(self.dates).isoformat()}
        elif self.texts:
            col.top = self.texts.most_common(TOP_VALUES)
        return col


def _table(name: str, rows, limits: Limits) -> Table | None:
    """rows：一列一列的值（第一個非空列當標題）。全空就回傳 None。"""
    header: list[str] | None = None
    stats: list[_ColumnStats] = []
    sample: list[list[str]] = []
    kept: list[list] = []  # 留著做分組小計（列數有上限）
    count = 0
    truncated = False
    for row in rows:
        values = list(row)
        if header is None:
            if not any(_cell(v) for v in values):
                continue
            while values and not _cell(values[-1]):
                values.pop()
            header = [_cell(v) or f"欄{i + 1}" for i, v in enumerate(values)]
            stats = [_ColumnStats(h) for h in header]
            continue
        if not any(_cell(v) for v in values):
            continue
        if count >= limits.max_rows:
            truncated = True
            break
        count += 1
        kept.append(values)
        for i, s in enumerate(stats):
            s.add(values[i] if i < len(values) else None)
        if len(sample) < SAMPLE_ROWS:
            sample.append([_cell(values[i]) if i < len(values) else "" for i in range(len(header))])
    if header is None:
        return None
    columns = [s.result() for s in stats]
    return Table(name=name, columns=columns, rows=count, sample=sample, truncated=truncated,
                 groups=_groups(columns, stats, kept))


def _groups(columns: list[Column], stats: list["_ColumnStats"], rows: list[list]) -> list[dict]:
    """用不同值不多的文字欄（部門、品項…）分組，算數字欄的合計與平均。"""
    nums = [i for i, c in enumerate(columns) if c.numeric][:GROUP_SUM_COLS]
    keys = [i for i, c in enumerate(columns)
            if c.top and 2 <= stats[i].distinct <= GROUP_MAX_KEYS and stats[i].filled >= len(rows) * 0.8
            and stats[i].distinct < stats[i].filled][:GROUP_BY_COLS]  # 每組只有一列就不用小計
    out = []
    for k in keys:
        acc: dict[str, list] = {}
        for r in rows:
            key = _cell(r[k]) if k < len(r) else ""
            if not key:
                continue
            g = acc.setdefault(key[:40], [0, [0.0] * len(nums), [0] * len(nums)])
            g[0] += 1
            for j, n in enumerate(nums):
                v = _number(r[n]) if n < len(r) else None
                if v is not None:
                    g[1][j] += v
                    g[2][j] += 1
        groups = [{"key": key, "count": c, "sums": [_round(x) for x in sums],
                   "means": [_round(x / k) if k else 0 for x, k in zip(sums, counts, strict=True)]}
                  for key, (c, sums, counts) in sorted(acc.items(), key=lambda kv: -kv[1][0])]
        out.append({"by": columns[k].name, "sums": [columns[n].name for n in nums], "groups": groups})
    return out


# ---------------------------------------------------------------- 各種格式


def _xlsx(doc: Document, data: bytes, limits: Limits) -> None:
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            unzipped = sum(i.file_size for i in z.infolist())
    except zipfile.BadZipFile:
        raise ReadError(f"「{doc.name}」不是有效的 xlsx 檔（可能是舊版 .xls 改了副檔名，或檔案損壞）") from None
    if unzipped > limits.max_unzipped_bytes:
        raise ReadError(f"「{doc.name}」解壓後有 {_size(unzipped)}，超過上限")
    try:
        # data_only：公式取 Excel 最後存檔時算好的值（沒有用 Excel 存過的檔案，公式格會是空的）
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # openpyxl 對壞檔的錯誤種類很多
        raise ReadError(f"「{doc.name}」讀不出來：{type(e).__name__}") from None
    try:
        sheets = wb.worksheets
        if len(sheets) > limits.max_sheets:
            doc.notes.append(f"共有 {len(sheets)} 個工作表，只讀前 {limits.max_sheets} 個")
        for ws in sheets[: limits.max_sheets]:
            t = _table(ws.title, ws.iter_rows(values_only=True), limits)
            if t is None:
                continue
            if t.truncated:
                doc.notes.append(f"工作表「{ws.title}」超過 {limits.max_rows:,} 列，後面沒有讀")
            doc.tables.append(t)
    finally:
        wb.close()
    if not doc.tables:
        raise ReadError(f"「{doc.name}」的工作表都是空的")


def _csv(doc: Document, data: bytes, limits: Limits) -> None:
    text, doc.encoding = decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    t = _table(PurePath(doc.name).stem, csv.reader(io.StringIO(text), dialect), limits)
    if t is None:
        raise ReadError(f"「{doc.name}」是空的")
    if t.truncated:
        doc.notes.append(f"超過 {limits.max_rows:,} 列，後面沒有讀")
    doc.tables.append(t)


def _txt(doc: Document, data: bytes, limits: Limits) -> None:
    text, doc.encoding = decode(data)
    doc.text = text.replace("\r\n", "\n").strip()
    if not doc.text:
        raise ReadError(f"「{doc.name}」是空的")


def _pdf(doc: Document, data: bytes, limits: Limits) -> None:
    import pdfplumber
    from pdfminer.pdfdocument import PDFPasswordIncorrect

    try:
        pdf = pdfplumber.open(io.BytesIO(data))
    except PDFPasswordIncorrect:
        raise ReadError(f"「{doc.name}」有密碼保護，請解除密碼後再上傳") from None
    except Exception as e:
        raise ReadError(f"「{doc.name}」不是有效的 PDF 或檔案損壞（{type(e).__name__}）") from None
    with pdf:
        doc.pages = len(pdf.pages)
        if doc.pages > limits.max_pages:
            doc.notes.append(f"共 {doc.pages} 頁，只讀前 {limits.max_pages} 頁")
        parts: list[str] = []
        images = 0
        for i, page in enumerate(pdf.pages[: limits.max_pages], 1):
            try:
                t = (page.extract_text() or "").strip()
            except Exception:
                t = ""
            images += len(page.images)
            if t:
                parts.append(f"〔第 {i} 頁〕\n{t}")
            page.close()  # 釋放這一頁的解析結果，大檔案才不會吃光記憶體
    doc.text = "\n\n".join(parts)
    read_pages = min(doc.pages, limits.max_pages)
    chars = len("".join(doc.text.split()))
    if chars < 20 * read_pages:  # 平均每頁不到 20 個字：大概是掃描檔
        if images:
            raise ReadError(f"「{doc.name}」看起來是掃描檔（整頁是圖片），讀不到文字；目前不支援掃描檔，請上傳有文字的 PDF")
        if not chars:
            raise ReadError(f"「{doc.name}」裡沒有文字")
