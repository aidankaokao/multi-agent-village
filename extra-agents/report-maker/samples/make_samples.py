"""產生測試用的三份範例檔（內容互相關聯：連鎖咖啡店 2025 上半年各分店營運）。

    DYLD_FALLBACK_LIBRARY_PATH=/opt/homebrew/lib uv run python samples/make_samples.py

- 分店營運數據.xlsx：每月數據表（分店 × 月份）＋ 分店基本資料
- 營運指標名詞解釋.txt：報表裡的指標定義、計算方式、公司目標值
- 上半年營運檢討報告.pdf：區經理的檢討報告（文字＋摘要表；數字取自 Excel，讓報告可以交叉比對）
"""

import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

OUT = Path(__file__).parent
random.seed(2025)

STORES = [  # 分店, 城市, 開幕年, 坪數, 員工數, 基準月營收, 成長趨勢
    ("台北信義店", "台北市", 2018, 62, 14, 1_850_000, 0.015),
    ("台北南港店", "台北市", 2021, 48, 10, 1_120_000, 0.035),
    ("桃園中壢店", "桃園市", 2019, 55, 11, 980_000, 0.005),
    ("台中逢甲店", "台中市", 2020, 45, 10, 1_050_000, 0.025),
    ("高雄巨蛋店", "高雄市", 2023, 58, 12, 890_000, -0.02),
]
MONTHS = ["2025-01", "2025-02", "2025-03", "2025-04", "2025-05", "2025-06"]
SEASON = [1.08, 0.92, 1.0, 0.98, 1.03, 1.06]  # 1 月年節、2 月淡季

rows = []
for name, city, opened, area, staff, base, trend in STORES:
    for i, m in enumerate(MONTHS):
        revenue = round(base * SEASON[i] * (1 + trend) ** i * random.uniform(0.97, 1.03), -2)
        ticket = round(random.uniform(158, 172) + (12 if city == "台北市" else 0))
        visitors = round(revenue / ticket)
        cogs_rate = random.uniform(0.31, 0.34) + (0.03 if name == "高雄巨蛋店" else 0)
        gross = round(revenue * (1 - cogs_rate), -2)
        labor = round(staff * 26_000 * random.uniform(0.97, 1.05), -2)  # 正職＋兼職平均
        rent = {"台北信義店": 260_000, "台北南港店": 175_000, "桃園中壢店": 120_000, "台中逢甲店": 110_000, "高雄巨蛋店": 135_000}[name]
        complaints = random.randint(1, 4) + (5 if name == "高雄巨蛋店" and i >= 3 else 0)
        nps = random.randint(38, 52) - (20 if name == "高雄巨蛋店" and i >= 3 else 0)
        rows.append([name, m, revenue, visitors, ticket, gross, labor, rent, complaints, nps])

# ---------------------------------------------------------------- Excel
wb = Workbook()
ws = wb.active
ws.title = "月營運數據"
header = ["分店", "月份", "營收", "來客數", "客單價", "毛利", "人事成本", "租金", "客訴件數", "NPS"]
ws.append(header)
for r in rows:
    ws.append(r)
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDE7F0")
for col, width in zip("ABCDEFGHIJ", [14, 10, 12, 10, 8, 12, 12, 10, 10, 8], strict=True):
    ws.column_dimensions[col].width = width
for row in ws.iter_rows(min_row=2, min_col=3, max_col=8):
    for c in row:
        c.number_format = "#,##0"

ws2 = wb.create_sheet("分店資料")
ws2.append(["分店", "城市", "開幕年", "坪數", "員工數", "店長"])
managers = ["林雅婷", "陳冠宇", "黃怡君", "張志豪", "李佩珊"]
for (name, city, opened, area, staff, *_), mgr in zip(STORES, managers, strict=True):
    ws2.append([name, city, opened, area, staff, mgr])
for c in ws2[1]:
    c.font = Font(bold=True)
wb.save(OUT / "分店營運數據.xlsx")

# ---------------------------------------------------------------- txt
(OUT / "營運指標名詞解釋.txt").write_text("""連鎖咖啡門市營運指標 名詞解釋（2025 年版）

本文件說明門市月報表中各項指標的定義與計算方式，以及公司 2025 年度的目標值，供店長與區經理對照使用。

【營收】
門市當月含稅銷售總額，扣除退貨與折讓。外送平台訂單以平台實際撥款金額計入。

【來客數】
當月結帳的交易筆數（一張收據算一位）。同一位顧客分兩次結帳算兩次。

【客單價】
營收 ÷ 來客數。反映顧客每次消費的平均金額。2025 年目標：台北市門市 175 元以上，其他縣市 160 元以上。

【毛利與毛利率】
毛利 = 營收 − 銷貨成本（咖啡豆、牛奶、食材、包材）。
毛利率 = 毛利 ÷ 營收。2025 年目標：65% 以上；低於 62% 須提出改善計畫。

【人事成本率】
人事成本 ÷ 營收。人事成本含正職薪資、兼職時薪、勞健保與加班費。2025 年目標：25% 以下。

【租金占比】
租金 ÷ 營收。超過 15% 視為高租金壓力門市，續約時需重新評估。

【坪效】
月營收 ÷ 門市坪數，衡量單位面積的營業效率。2025 年目標：每坪每月 20,000 元以上。

【客訴件數】
當月經客服中心或門市登記的正式客訴（不含一般意見回饋）。單月 5 件以上須於週會說明原因。

【NPS（淨推薦值）】
問卷題目「您有多大可能推薦本店給朋友？」0～10 分。
NPS = 推薦者（9～10 分）比例 − 批評者（0～6 分）比例，範圍 −100～100。
2025 年目標：40 以上；低於 30 列為服務品質觀察門市。

【營業利益（門市層級）】
毛利 − 人事成本 − 租金 − 其他營業費用（水電、耗材、行銷分攤）。月報表未列其他營業費用，概估約為營收的 8%。
""", encoding="utf-8")

# ---------------------------------------------------------------- PDF（數字取自上面的資料）
def total(store=None, months=None, idx=2):
    return sum(r[idx] for r in rows if (store is None or r[0] == store) and (months is None or r[1] in months))


Q1, Q2 = MONTHS[:3], MONTHS[3:]
summary = []
for name, city, opened, area, staff, *_ in STORES:
    rev = total(name)
    gm = total(name, idx=5) / rev
    labor = total(name, idx=6) / rev
    growth = total(name, Q2) / total(name, Q1) - 1
    summary.append((name, rev, gm, labor, growth, rev / 6 / area))

best = max(summary, key=lambda s: s[4])
worst = min(summary, key=lambda s: s[4])
kh = [r for r in rows if r[0] == "高雄巨蛋店"]

rows_html = "".join(
    f"<tr><td>{n}</td><td class=n>{rev:,.0f}</td><td class=n>{gm:.1%}</td><td class=n>{lab:.1%}</td>"
    f"<td class=n>{g:+.1%}</td><td class=n>{eff:,.0f}</td></tr>"
    for n, rev, gm, lab, g, eff in summary
)
html = f"""<!doctype html><html lang="zh-Hant"><head><meta charset="utf-8"><style>
@page {{ size: A4; margin: 22mm 20mm; @bottom-center {{ content: "第 " counter(page) " 頁"; font-size: 8pt; color: #888; }} }}
html {{ font-family: "Noto Sans CJK TC", "PingFang TC", sans-serif; font-size: 10.5pt; line-height: 1.75; color: #222; }}
h1 {{ font-size: 18pt; margin: 0; }} h2 {{ font-size: 12.5pt; margin: 16pt 0 6pt; border-bottom: 1pt solid #999; }}
.meta {{ color: #666; font-size: 9pt; margin-bottom: 10pt; }}
table {{ border-collapse: collapse; width: 100%; font-size: 9pt; }} th, td {{ border: .6pt solid #aaa; padding: 3pt 5pt; }}
th {{ background: #eee; }} td.n {{ text-align: right; }}
</style></head><body>
<h1>2025 年上半年 門市營運檢討報告</h1>
<div class="meta">報告人：北中南區域經理 王建明　｜　日期：2025 年 7 月 8 日　｜　範圍：5 家直營門市，2025 年 1～6 月</div>

<h2>一、整體概況</h2>
<p>上半年 5 家門市合計營收 {total():,.0f} 元，其中第二季營收 {total(months=Q2):,.0f} 元，較第一季的 {total(months=Q1):,.0f} 元
{"成長" if total(months=Q2) >= total(months=Q1) else "衰退"} {abs(total(months=Q2) / total(months=Q1) - 1):.1%}。
1 月受惠於春節檔期與年節禮盒，2 月則為傳統淡季，營收普遍下滑；4 月起隨天氣轉熱，冰飲比例提高，營收回溫。</p>
<p>整體毛利率維持在公司目標 65% 附近，但各店差異明顯；人事成本受基本工資調漲影響，多數門市較去年同期上升約 1～2 個百分點。</p>

<h2>二、各門市上半年摘要</h2>
<table><tr><th>門市</th><th>上半年營收（元）</th><th>毛利率</th><th>人事成本率</th><th>Q2 較 Q1 營收</th><th>月平均坪效（元）</th></tr>
{rows_html}</table>

<h2>三、重點觀察</h2>
<p><b>1. {best[0]}成長最快。</b>第二季營收較第一季成長 {best[4]:.1%}，主因是周邊辦公大樓進駐率提高，平日早晨外帶訂單明顯增加。
建議下半年評估延長早班人力，並與周邊企業洽談團購方案。</p>
<p><b>2. 高雄巨蛋店需要關注。</b>第二季營收較第一季變化 {worst[4]:+.1%}，為五店最差。4 月起更換主要烘焙供應商後，
顧客反映餐點品質不穩定：4～6 月客訴分別為 {kh[3][8]}、{kh[4][8]}、{kh[5][8]} 件，NPS 分別為 {kh[3][9]}、{kh[4][9]}、{kh[5][9]}，
已低於公司「服務品質觀察門市」的標準。另因食材成本較高，毛利率 {[s for s in summary if s[0] == "高雄巨蛋店"][0][2]:.1%} 也是五店最低。</p>
<p><b>3. 台北信義店營收最高但租金壓力大。</b>每月租金 260,000 元，續約時程在 2026 年初，需提前評估租金占比是否超過 15% 的警戒線。</p>
<p><b>4. 桃園中壢店表現平穩。</b>營收成長有限，但客訴少、NPS 穩定，可作為標準作業流程的示範門市。</p>

<h2>四、下半年行動計畫</h2>
<ul>
<li>高雄巨蛋店：8 月底前完成供應商重新評選，9 月起每週回報客訴與 NPS。</li>
<li>台北南港店：試行早班加開一個收銀台，觀察尖峰時段等候時間。</li>
<li>全區：導入排班系統，以人事成本率 25% 為目標，避免尖離峰人力錯置。</li>
<li>台北信義店：委請總部展店部門於 10 月前提供周邊租金行情與續約建議。</li>
</ul>
<p style="color:#666;font-size:9pt">附註：指標定義請參閱《營運指標名詞解釋》；詳細月數據請參閱《分店營運數據》Excel 檔。</p>
</body></html>"""

from weasyprint import HTML  # noqa: E402

HTML(string=html).write_pdf(OUT / "上半年營運檢討報告.pdf")
print("已產生：", *sorted(p.name for p in OUT.iterdir() if p.suffix in (".xlsx", ".txt", ".pdf")))
